"""ForEvento: geração de certificados em Python, sem Google Apps Script.

Usa o modelo persistente do evento, os registros SQLite existentes e não
altera esquemas nem sobrescreve certificados previamente gerados.
"""
from pathlib import Path
import datetime
import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid

def _pdf_from_pptx(source, destination):
    """Conversão isolada: soffice.com no Windows, LibreOffice headless no Linux."""
    import urllib.parse
    candidates=[]
    configured=os.environ.get('FOREVENTO_LIBREOFFICE','').strip()
    if configured:candidates.append(configured)
    if os.name=='nt':
        for root in (os.environ.get('PROGRAMFILES',r'C:\Program Files'),os.environ.get('PROGRAMFILES(X86)',r'C:\Program Files (x86)')):
            candidates.extend([str(Path(root)/'LibreOffice/program/soffice.com'),str(Path(root)/'LibreOffice/program/soffice.exe')])
        candidates.extend(['soffice.com','soffice.exe'])
    else:candidates.extend(['libreoffice','soffice'])
    exe=next((str(Path(x)) if Path(x).is_file() else shutil.which(x) for x in candidates if Path(x).is_file() or shutil.which(x)),None)
    if not exe:raise RuntimeError('LibreOffice não encontrado. Configure FOREVENTO_LIBREOFFICE com o caminho do soffice.com.')
    with tempfile.TemporaryDirectory(prefix='forevento_lo_') as work:
        workpath=Path(work)
        profile=(workpath/'perfil').resolve().as_uri()
        proc=subprocess.run([exe,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf:impress_pdf_Export','--outdir',str(workpath),str(Path(source).resolve())],stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=120,check=False)
        produced=workpath/(Path(source).stem+'.pdf')
        if proc.returncode!=0 or not produced.is_file() or produced.stat().st_size<100:
            out=(proc.stdout+proc.stderr).decode('utf-8',errors='replace')[-1200:]
            raise RuntimeError('Falha ao converter PPTX com LibreOffice (código '+str(proc.returncode)+'): '+out)
        shutil.copyfile(produced,destination)

def _assinatura_imagem(modelos_dir=None):
    """Localiza a assinatura existente; se necessário, extrai uma imagem de modelo PDF automático.

    Nunca confunde o QR ou logotipo com a assinatura: exige um único candidato
    gráfico no quadrante inferior esquerdo. Não modifica o PDF de origem.
    """
    configured=os.environ.get("FOREVENTO_ASSINATURA_IMAGEM", "").strip()
    app_dir=Path(__file__).resolve().parent
    roots=(app_dir/"static"/"images",app_dir/"static",app_dir/"assets",app_dir.parent/"storage",app_dir.parent/"uploads")
    candidates=([Path(configured).expanduser()] if configured else [])
    for root in roots:
        for name in ("assinatura.png","assinatura.jpg","assinatura.jpeg","assinatura_autoridade.png","assinatura_autoridade.jpg"):
            candidates.append(root/name)
    for path in candidates:
        if path.is_file() and path.suffix.lower() in (".png",".jpg",".jpeg"):
            return path
    if modelos_dir is not None:
        import fitz
        from PIL import Image as PILImage
        directory=Path(modelos_dir).resolve()
        pdfs=sorted(p for p in directory.glob("*.pdf") if p.is_file() and re.search(r"_a\.pdf$",p.name,re.I))
        # The source is only a candidate if its graphic is independently embedded.
        # A flattened PDF or ambiguous image positions must not be guessed.
        found=[]
        for pdf in pdfs:
            with fitz.open(str(pdf)) as doc:
                if not doc or doc.is_encrypted:continue
                page=doc[0]
                for info in page.get_image_info(xrefs=True):
                    bbox=fitz.Rect(info['bbox'])
                    if bbox.is_empty or not info.get('xref'):continue
                    if bbox.x0 >= page.rect.width*0.50 or bbox.y0 < page.rect.height*0.60:continue
                    if bbox.width < 45 or bbox.height < 10 or bbox.width/bbox.height < 1.3:continue
                    found.append((pdf,int(info['xref'])))
        unique={(str(p),xref) for p,xref in found}
        if len(unique)==1:
            pdf,xref=found[0]
            with fitz.open(str(pdf)) as doc:
                raw=doc.extract_image(xref)['image']
            import io
            with PILImage.open(io.BytesIO(raw)) as im:
                rgba=im.convert('RGBA')
                cache=app_dir.parent/'storage'/'assinatura_extraida_pdf.png'
                cache.parent.mkdir(parents=True,exist_ok=True)
                rgba.save(cache)
            return cache
        if len(unique)>1:
            raise RuntimeError('Assinatura não identificada com segurança: múltiplas imagens na região inferior esquerda dos modelos PDF. Não foi alterado nenhum certificado.')
    raise RuntimeError('Assinatura não localizada: nenhum arquivo de assinatura nem imagem independente identificável em modelo PDF automático. Não é seguro usar o logotipo ou o QR Code.')

def _replace_text_in_pptx(src,dst,fields,qr_bytes=None,modelos_dir=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs=Presentation(str(src))
    signature_tags=[]
    def visit(shapes):
        for shape in shapes:
            if getattr(shape,"has_text_frame",False):
                for p in shape.text_frame.paragraphs:
                    full="".join(run.text for run in p.runs)
                    changed=full
                    if re.search(r"<<\s*assinatura\s*>>|\{\{\s*assinatura\s*\}\}",full,flags=re.I):
                        signature_tags.append((shape, p))
                        changed=re.sub(r"<<\s*assinatura\s*>>|\{\{\s*assinatura\s*\}\}","",changed,flags=re.I)
                    for k,v in fields.items():
                        changed=re.sub(r"<<\s*"+re.escape(k)+r"\s*>>|\{\{\s*"+re.escape(k)+r"\s*\}\}",
                                       lambda m:str(v),changed,flags=re.I)
                    if changed!=full:
                        if p.runs:
                            p.runs[0].text=changed
                            for r in p.runs[1:]:r.text=""
                        else:p.text=changed
            if getattr(shape,"has_table",False):
                for row in shape.table.rows:
                    for cell in row.cells:visit([cell])
    for slide in prs.slides:
        visit(slide.shapes)
        if qr_bytes:
            slide.shapes.add_picture(io.BytesIO(qr_bytes),
                prs.slide_width-Inches(1.25),prs.slide_height-Inches(1.25),
                width=Inches(0.9),height=Inches(0.9))
    if signature_tags:
        image_path=_assinatura_imagem(modelos_dir)
        for shape, _paragraph in signature_tags:
            # Coordenadas da TAG no modelo, com o ajuste solicitado aplicado somente à assinatura.
            slide=next((sl for sl in prs.slides if any(sh.shape_id == shape.shape_id and sh._element is shape._element for sh in sl.shapes)),None)
            if slide is None:
                raise RuntimeError("TAG de assinatura em estrutura PPTX não suportada; use uma caixa de texto independente.")
            slide.shapes.add_picture(str(image_path),shape.left+Pt(8.50),shape.top-Pt(14.17),
                                     width=Inches(1.4))
    prs.save(str(dst))

def _qr_png(url):
    import qrcode
    b=io.BytesIO()
    qrcode.make(url).save(b,format="PNG")
    return b.getvalue()

CREDENCIAL_DX_PT = 3 * 72 / 25.4  # 3 mm para a direita
CREDENCIAL_DY_PT = 5 * 72 / 25.4  # 5 mm para cima
URL_ABAIXO_CODIGO_PT = 11  # distância vertical entre as linhas, em pontos

def _append_validation_footer(pdf_path, fields):
    """Mesmo rodapé automático usado para modelos PDF, sem TAG no PPTX."""
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()
    for page in reader.pages:
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        overlay = io.BytesIO()
        c = canvas.Canvas(overlay, pagesize=(width, height))
        c.setFont("Helvetica", 8)
        c.drawString(24 + CREDENCIAL_DX_PT, 22 + CREDENCIAL_DY_PT, "Código: " + str(fields["codigovalidacao"]))
        c.drawString(24 + CREDENCIAL_DX_PT, 22 + CREDENCIAL_DY_PT - URL_ABAIXO_CODIGO_PT, str(fields["urlvalidacao"])[:120])
        c.save()
        overlay.seek(0)
        page.merge_page(PdfReader(overlay).pages[0])
        writer.add_page(page)
    with tempfile.NamedTemporaryFile(suffix=".pdf", dir=str(Path(pdf_path).parent), delete=False) as f:
        temp = Path(f.name)
    try:
        with temp.open("wb") as f:
            writer.write(f)
        os.replace(temp, pdf_path)
    finally:
        temp.unlink(missing_ok=True)

def _render_pdf(source, output, fields, qr_bytes, automatic):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from pypdf import PdfReader,PdfWriter
    from PIL import Image
    ext=source.suffix.lower()
    if ext==".pptx":
        with tempfile.TemporaryDirectory() as td:
            tmp=Path(td)
            prepared=tmp/"prepared.pptx"
            _replace_text_in_pptx(source,prepared,fields,qr_bytes if automatic else None,source.parent) if automatic else shutil.copyfile(source,prepared)
            _pdf_from_pptx(prepared,tmp/"prepared.pdf")
            shutil.copyfile(tmp/"prepared.pdf",output)
            if automatic:
                _append_validation_footer(output, fields)
        return
    if ext==".pdf":
        if not automatic:
            shutil.copyfile(source,output)
            return
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("Instale PyMuPDF: python -m pip install pymupdf") from exc
        from reportlab.lib.utils import ImageReader
        from pypdf import PdfReader,PdfWriter
        aliases={
            "nome":["Nome","NOME","NOME_PARTICIPANTE"],
            "cpf":["CPF"],
            "atvdata":["atvData","DATA_ATIVIDADE"],
            "tipoinscricao":["tipoInscricao","TIPO_INSCRICAO"],
            "cargahoraria":["cargaHoraria","CARGA_HORARIA"],
            "evento":["Evento","EVENTO","NOME_EVENTO"],
            "atividade":["Atividade","ATIVIDADE","NOME_ATIVIDADE"],
            "tituloatividade":["TituloAtividade","TITULO_ATIVIDADE","TITULOATIVIDADE"],
            "email":["Email","EMAIL"],
            "codigovalidacao":["CodigoValidacao","CODIGO_VALIDACAO"],
            "urlvalidacao":["UrlValidacao","URL_VALIDACAO"],
            "codigo":["Codigo","CODIGO"],
            "url":["Url","URL"],
            "codigovalidacaoextra":["CODIGO_CREDENCIAL"],
        }
        doc=fitz.open(str(source))
        if doc.is_encrypted:
            doc.close()
            raise RuntimeError("O modelo PDF está protegido por senha.")
        found=set()
        try:
            for page in doc:
                changes=[]
                for key,names in aliases.items():
                    for name in names:
                        for tag in ("{{"+name+"}}","<<"+name+">>"):
                            for rect in page.search_for(tag):
                                changes.append((rect,str(fields.get(key,"")),key))
                                found.add(key)
                signature_rects=[]
                for signature_tag in ("<<assinatura>>", "{{assinatura}}", "<<ASSINATURA>>", "{{ASSINATURA}}"):
                    signature_rects.extend(page.search_for(signature_tag))
                for rect in signature_rects:
                    page.add_redact_annot(rect,fill=(1,1,1))
                qr_rects=[]
                for qr_tag in ("<<QR_CODE>>", "{{QR_CODE}}"):
                    qr_rects.extend(page.search_for(qr_tag))
                for rqr in qr_rects:
                    page.add_redact_annot(rqr,fill=(1,1,1))
                for rect,_,_ in changes:
                    page.add_redact_annot(rect,fill=(1,1,1))
                if changes or qr_rects or signature_rects:page.apply_redactions(images=0,graphics=0)
                if signature_rects:
                    signature_image=_assinatura_imagem(source.parent)
                    for r in signature_rects:
                        from PIL import Image as PILImage
                        with PILImage.open(signature_image) as im:
                            ratio=im.height/im.width
                        width=100.8  # 1.4 polegadas, igual ao PPTX
                        x=r.x0+8.50; y=r.y0-14.17
                        page.insert_image(fitz.Rect(x,y,x+width,y+width*ratio),filename=str(signature_image),overlay=True,keep_proportion=True)

                for rect,value,key in changes:
                    if not value:continue
                    size=max(7,min(18,rect.height*0.85))
                    # The original placeholder must leave enough space for the real value.
                    area=fitz.Rect(rect.x0,rect.y0-2,
                                   min(page.rect.width-16,max(rect.x1,rect.x0+page.rect.width*0.55)),
                                   rect.y1+max(12,rect.height))
                    result=-1
                    while size>=7 and result<0:
                        result=page.insert_textbox(area,value,fontsize=size,
                             fontname="helv",color=(0,0,0),align=0)
                        size-=0.5
                    if result<0:
                        raise RuntimeError("Texto extenso demais para o campo "+key+
                            ". Aumente a área reservada no modelo PDF.")
                if qr_bytes and qr_rects:
                    r=qr_rects[0]
                    side=min(85,page.rect.height-r.y0-8,page.rect.width-r.x0-8)
                    if side>20:page.insert_image(fitz.Rect(r.x0,r.y0,r.x0+side,r.y0+side),stream=qr_bytes,keep_proportion=True,overlay=True)
                # TAGs não reconhecidas nunca devem ser distribuídas como certificado final.
                leftovers=[]
                for block in page.get_text("blocks"):
                    leftovers.extend(re.findall(r'<<[^<>]{2,60}>>|\{\{[^{}]{2,60}\}\}',str(block[4])))
                if leftovers:
                    raise RuntimeError("TAGs não preenchidas no PDF: "+", ".join(sorted(set(leftovers))[:12]))
            if not found.intersection({"nome","evento","atividade","tituloatividade"}):
                raise RuntimeError("PDF sem TAGs de texto pesquisáveis. Exporte o modelo mantendo as TAGs como texto selecionável, por exemplo {{Nome}} e {{Evento}}. PDFs achatados/imagens precisam de campos posicionados.")
            with tempfile.TemporaryDirectory() as td:
                prepared=Path(td)/"filled.pdf"
                doc.save(str(prepared),garbage=4,deflate=True)
                reader=PdfReader(str(prepared));writer=PdfWriter()
                for page in reader.pages:
                    width,height=float(page.mediabox.width),float(page.mediabox.height)
                    overlay=io.BytesIO()
                    c=canvas.Canvas(overlay,pagesize=(width,height))
                    c.setFont("Helvetica",8)
                    if "codigovalidacao" not in found:
                        c.drawString(24+CREDENCIAL_DX_PT,22+CREDENCIAL_DY_PT,"Código: "+fields["codigovalidacao"])
                    if "urlvalidacao" not in found:
                        c.drawString(24+CREDENCIAL_DX_PT,22+CREDENCIAL_DY_PT-URL_ABAIXO_CODIGO_PT,fields["urlvalidacao"][:120])
                    if qr_bytes and not qr_rects:c.drawImage(ImageReader(io.BytesIO(qr_bytes)),width-92,12,72,72)
                    c.save();overlay.seek(0)
                    page.merge_page(PdfReader(overlay).pages[0]);writer.add_page(page)
                with open(output,"wb") as f:writer.write(f)
        finally:
            doc.close()
        return
    if ext in (".png",".jpg",".jpeg"):
        from reportlab.lib.utils import ImageReader
        with Image.open(source) as im:w,h=im.size
        c=canvas.Canvas(str(output),pagesize=(w,h))
        c.drawImage(ImageReader(str(source)),0,0,w,h)
        if automatic:
            c.setFont("Helvetica",max(9,min(16,w/85)))
            c.drawString(18+CREDENCIAL_DX_PT,26+CREDENCIAL_DY_PT,"Código: "+fields["codigovalidacao"])
            c.drawString(18+CREDENCIAL_DX_PT,26+CREDENCIAL_DY_PT-URL_ABAIXO_CODIGO_PT,fields["urlvalidacao"][:120])
            if qr_bytes:c.drawImage(ImageReader(io.BytesIO(qr_bytes)),w-110,15,90,90)
        c.save();return
    raise RuntimeError("Formato de modelo não suportado: "+ext)

def gerar_em_massa(db_connect, modelos_dir, base_dir, event_id, base_url, model_id="", model_format=""):
    import json
    model_format=str(model_format or "").lower().strip()
    if model_format not in ("", "pdf", "pptx"):
        return {"sucesso":False,"mensagem":"Formato de modelo inválido."}
    event_id=str(event_id or "").strip()
    if not event_id:return {"sucesso":False,"mensagem":"Selecione um evento."}
    c=db_connect()
    created_files=[]
    try:
        event=c.execute("select id,name from events where id=?",(event_id,)).fetchone()
        if not event:return {"sucesso":False,"mensagem":"Evento não encontrado."}
        setting=c.execute("select v from settings where k=?",("certcfg:"+event_id,)).fetchone()
        cfg=json.loads(setting["v"]) if setting and setting["v"] else {}
        available=c.execute("select * from certificate_models order by id").fetchall()
        managed=Path(modelos_dir).resolve()
        available=[m for m in available if Path(m["file_path"]).resolve().parent==managed
                   and Path(m["file_path"]).is_file() and Path(m["file_path"]).suffix.lower() in (".pdf", ".pptx")]
        if model_format:
            available=[m for m in available if Path(m["file_path"]).suffix.lower()=="."+model_format]
        if not available:
            return {"sucesso":False,"mensagem":"Nenhum modelo PDF ou PPTX cadastrado no servidor com sufixo _a ou _m."}
        # Seleção explícita: nunca trocar silenciosamente um formato por outro.
        chosen=next((m for m in available if str(m["id"])==str(model_id)),None) if model_id else None
        if model_id and chosen is None:
            return {"sucesso":False,"mensagem":"Modelo selecionado não encontrado no armazenamento do servidor."}
        signature=str(cfg.get("tipoAssinatura") or "AUTOMATICA").upper()
        if signature not in ("AUTOMATICA","MANUAL"):signature="AUTOMATICA"
        regs=c.execute("""select r.*,u.name nome,u.language idioma
          from registrations r left join users u on u.id=r.user_id
          where r.event_id=? and upper(coalesce(r.status,'')) not in
          ('CANCELADO','CANCELADA','REEMBOLSADO') order by r.created_at,r.id""",(event_id,)).fetchall()
        # Conferência de presença e pagamento ocorre no servidor.
        from .main import _inscricao_liberada_para_presenca_certificado
        generated=0;ignored=0;details=[]
        storage=(Path(base_dir)/"storage").resolve()
        if storage.exists() and not storage.is_dir():
            raise RuntimeError("O caminho storage existe como ARQUIVO, não como pasta. Preserve-o e verifique o conflito antes de gerar.")
        dest=(storage/"certificados").resolve()
        if dest.exists() and not dest.is_dir():
            raise RuntimeError("O caminho storage/certificados existe como ARQUIVO, não como pasta.")
        for reg in regs:
            rid=reg["id"]
            idioma=str(reg["idioma"] or "pt").strip().lower().replace("-","_")
            idioma=idioma.split("_")[0]
            # O seletor de conferência NÃO determina o modelo de geração.
            matches=[m for m in available if str(m["language"] or "").lower()[:2]==idioma[:2]
                     and str(m["signature_type"] or "").upper()==signature]
            if chosen:
                model=chosen if str(chosen["language"] or "").lower()[:2]==idioma[:2] and str(chosen["signature_type"] or "").upper()==signature else None
            else:
                model=matches[0] if len(matches)==1 else None
            if model is None:
                ignored+=1
                details.append(str(reg["email"] or rid)+": falta modelo compatível ou há mais de um modelo para idioma "+idioma+" e assinatura "+signature+"; selecione explicitamente o modelo")
                continue
            source=Path(model["file_path"]).resolve()
            automatic=str(model["signature_type"] or "").upper()=="AUTOMATICA"
            if not _inscricao_liberada_para_presenca_certificado(c,reg):
                ignored+=1;continue
            pres=c.execute("""select status from attendance where event_id=? and activity=?
              and (user_id=? or lower(email)=lower(?))
              order by id desc limit 1""",(event_id,reg["activity"],reg["user_id"],reg["email"])).fetchone()
            if not pres or str(pres["status"] or "").upper().strip() not in ("PRESENTE","PRESENÇA","PRESENCA","CONFIRMADO","CONFIRMADA"):
                ignored+=1;continue
            old=c.execute("""select 1 from certificates where event_id=? and activity=?
              and (user_id=? or lower(email)=lower(?)) limit 1""",
              (event_id,reg["activity"],reg["user_id"],reg["email"])).fetchone()
            if old:ignored+=1;continue
            cid="CERT-"+uuid.uuid4().hex[:16].upper()
            code="FE-"+uuid.uuid4().hex[:12].upper()
            validation=base_url.rstrip("/")+"/validar-certificado?codigo="+code
            from .main import titulo_atividade_para_exibicao
            titulo_atividade=titulo_atividade_para_exibicao(c,event_id,reg["activity"])
            schedule=c.execute("""select date,workload from schedules where event_id=?
              and (titulo=? or synopsis_activity=? or synopsis=?) order by id limit 1""",
              (event_id,titulo_atividade,reg["activity"],reg["activity"])).fetchone()
            person=c.execute("select cpf from users where id=?",(reg["user_id"],)).fetchone() if reg["user_id"] else None
            fields={
              "cpf":(person["cpf"] if person else "") or "",
              "atvdata":(schedule["date"] if schedule else "") or "",
              "tipoinscricao":reg["type"] or "",
              "cargahoraria":(schedule["workload"] if schedule else "") or "",
              "codigovalidacaoextra":code,
              "nome":reg["nome"] or reg["email"] or "",
              "evento":event["name"],"atividade":titulo_atividade,
              "tituloatividade":titulo_atividade,
              "email":reg["email"] or "",
              "codigovalidacao":code,"urlvalidacao":validation,
              "codigo":code,"url":validation,
            }
            qr=_qr_png(validation) if automatic else None
            dest.mkdir(parents=True,exist_ok=True)
            target=dest/(cid+".pdf")
            try:
                _render_pdf(source,target,fields,qr,automatic)
                if not target.is_file() or target.stat().st_size<100:
                    raise RuntimeError("PDF gerado está vazio.")
                created_files.append(target)
                # ID aleatório impede colisões; publicação só após envio.
                c.execute("""insert into certificates
                  (id,event_id,user_id,email,activity,status,validation_code,url,file_url,
                   signature_type,language,created_at)
                  values(?,?,?,?,?,?,?,?,?,?,?,?)""",
                  (cid,event_id,reg["user_id"],reg["email"],reg["activity"],
                   "GERADO",code,validation,
                   "/api/certificados/arquivo/"+cid,
                   "AUTOMATICA" if automatic else "MANUAL",
                   model["language"],datetime.datetime.now().isoformat(timespec="seconds")))
                generated+=1
            except Exception as exc:
                target.unlink(missing_ok=True)
                details.append(str(reg["email"] or rid)+": "+str(exc))
        c.commit()
        return {"sucesso":generated>0,
                "mensagem":f"{generated} certificado(s) gerado(s); {ignored} ignorado(s)."+
                    (" Ocorreram erros em alguns registros." if details else ""),
                "gerados":generated,"ignorados":ignored,"detalhes":details}
    except Exception as exc:
        c.rollback()
        for p in created_files:p.unlink(missing_ok=True)
        return {"sucesso":False,"mensagem":"Erro na geração: "+str(exc)}
    finally:c.close()
